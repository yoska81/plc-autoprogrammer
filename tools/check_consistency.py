"""Cross-consistency between deliverables of one variant:
part numbers / quantities (BOM CSV, Excel BOM, drawings, engineering package), wire numbers
(wire list CSV, Excel, labels, drawings), terminals (terminal plan CSV, Excel, labels),
I/O addresses (I/O CSV, Excel, nicknames CSV, drawings, package), alarm texts (HMI, Excel,
package), revision letter (all title blocks / footers), motor FLA against an independent copy of
NEC 430.250 / 430.248 (CEC Table 44/45 use the same values), and that the 24 VDC load and heat
calculations add up. Usage: tools/.venv/bin/python tools/check_consistency.py [variant ...]
"""
import csv
import io
import json
import re
import sys
from pathlib import Path

import openpyxl
from pypdf import PdfReader

OUT = Path(__file__).parent / 'out'
results = []

# NEC Table 430.250 (3-ph, 460/230/208/575 V columns) and 430.248 (1-ph 115 V) — typed in independently.
NEC_3PH = {0.5: (1.1, 2.2, 2.4, 0.9), 0.75: (1.6, 3.2, 3.5, 1.3), 1: (2.1, 4.2, 4.6, 1.7), 1.5: (3.0, 6.0, 6.6, 2.4), 2: (3.4, 6.8, 7.5, 2.7),
           3: (4.8, 9.6, 10.6, 3.9), 5: (7.6, 15.2, 16.7, 6.1), 7.5: (11, 22, 24.2, 9), 10: (14, 28, 30.8, 11), 15: (21, 42, 46.2, 17),
           20: (27, 54, 59.4, 22), 25: (34, 68, 74.8, 27), 30: (40, 80, 88, 32), 40: (52, 104, 114, 41), 50: (65, 130, 143, 52),
           60: (77, 154, 169, 62), 75: (96, 192, 211, 77), 100: (124, 248, 273, 99)}
NEC_1PH_115 = {1 / 6: 4.4, 0.25: 5.8, 1 / 3: 7.2, 0.5: 9.8, 0.75: 13.8, 1: 16, 1.5: 20, 2: 24, 3: 34, 5: 56}
COL = {'480VAC-3PH': 0, '240VAC-3PH': 1, '208VAC-3PH': 2, '600VAC-3PH': 3}


def report(level, where, msg):
    results.append((level, where, msg))


def read_csv(p: Path):
    return list(csv.DictReader(io.StringIO(p.read_text(encoding='utf-8-sig'))))


def pdf_text(p: Path):
    return [pg.extract_text() or '' for pg in PdfReader(p).pages]


def sheet_rows(wb, name):
    ws = wb[name]
    rows = [r for r in ws.iter_rows(values_only=True)]
    hdr = [str(c) if c is not None else '' for c in rows[0]]
    return [dict(zip(hdr, r)) for r in rows[1:] if any(c not in (None, '') for c in r)]


# IEC 60947-4-1 Annex G rated currents (kW, 230 V, 400 V, 690 V), next rating up — 50 Hz supplies.
IEC_MOTOR = [(0.37, 1.9, 1.1, 0.64), (0.55, 2.6, 1.5, 0.87), (0.75, 3.3, 1.9, 1.1), (1.1, 4.7, 2.7, 1.6), (1.5, 6.3, 3.6, 2.1),
             (2.2, 8.5, 4.9, 2.8), (3, 11.3, 6.5, 3.8), (4, 15, 8.5, 4.9), (5.5, 20, 11.5, 6.7), (7.5, 27, 15.5, 9), (11, 39, 22, 12.9)]


def nec_fla(hp, volt):
    if volt in ('400VAC-3PH', '690VAC-3PH'):
        kw = hp * 0.7457
        row = next((r for r in IEC_MOTOR if r[0] >= kw - 0.005), IEC_MOTOR[-1])
        return row[3] if volt == '690VAC-3PH' else row[2]
    if volt == '120VAC-1PH':
        k = min((h for h in NEC_1PH_115 if h >= hp - 1e-9), default=5)
        return NEC_1PH_115[k]
    k = min((h for h in NEC_3PH if h >= hp - 1e-9), default=100)
    return NEC_3PH[k][COL[volt]]


def check(vdir: Path):
    v = vdir.name
    z = vdir / 'zip'
    data = json.loads((vdir / 'data.json').read_text())
    base = data['fileBase']
    wb = openpyxl.load_workbook(next(z.glob('*_report.xlsx')))
    drw = pdf_text(z / 'electrical' / f'{base}_drawings.pdf')
    eng = pdf_text(z / f'{base}_engineering_package.pdf')
    drw_all = '\n'.join(drw)
    eng_all = '\n'.join(eng)
    flat = lambda s: re.sub(r'\s+', ' ', s)  # noqa: E731
    drw_flat, eng_flat = flat(drw_all), flat(eng_all)

    # ---- BOM
    bom_csv = read_csv(z / 'electrical' / 'bom.csv')
    csv_q = {}
    for r in bom_csv:
        csv_q[r['Part Number']] = csv_q.get(r['Part Number'], 0) + int(r['Qty'])
    csv_q.pop('', None)
    xl_q = {}
    for r in sheet_rows(wb, 'BOM'):
        if r.get('Part number') and isinstance(r.get('Qty'), (int, float)):
            xl_q[str(r['Part number'])] = xl_q.get(str(r['Part number']), 0) + r['Qty']
    data_q = {}
    for g in data['bom']:
        for ln in g['lines']:
            if ln.get('partNumber'):
                data_q[ln['partNumber']] = data_q.get(ln['partNumber'], 0) + ln['qty']
    if csv_q != xl_q:
        report('FAIL', v, f'BOM CSV vs Excel differ: {sorted(set(csv_q.items()) ^ set(xl_q.items()))[:6]}')
    if {k: v_ for k, v_ in csv_q.items() if k} != {k: v_ for k, v_ in data_q.items() if k and k in csv_q}:
        report('FAIL', v, 'BOM CSV vs derived BOM differ')
    for pn in csv_q:
        if pn and pn not in drw_flat:
            report('FAIL', v, f'part {pn} (BOM) not found anywhere in the drawing set')
        if pn and pn not in eng_flat:
            report('FAIL', v, f'part {pn} (BOM) not found in the engineering package')
    # device tags: labels vs BOM tags
    labels = read_csv(vdir / 'labels.csv')
    bom_tags = set()
    for r in bom_csv:
        bom_tags.update(t.strip() for t in re.split(r'[,\s]+', r['Device Tags']) if t.strip())
    for l in labels:
        if l['Type'] == 'Cabinet device tags' and l['Number'] not in bom_tags:
            report('FAIL', v, f'device label {l["Number"]} has no BOM line')

    # ---- Label PDFs: one label-printer page per label copy; ZIP labels CSV = labels data
    lp = next((z / 'labels').glob('*_label_printer.pdf'), None) if (z / 'labels').exists() else None
    zl = next((z / 'labels').glob('*_labels.csv'), None) if (z / 'labels').exists() else None
    if zl:
        zrows = read_csv(zl)
        copies = sum(int(r.get('Copies') or 1) for r in zrows)
        if {(r['Type'], r['Number']) for r in zrows} != {(r['Type'], r['Number']) for r in labels}:
            report('FAIL', v, 'ZIP labels CSV differs from the labels data')
        if lp and len(PdfReader(lp).pages) != copies:
            report('FAIL', v, f'label printer PDF has {len(PdfReader(lp).pages)} labels, CSV copies total {copies}')
    # ---- Wires
    wl = read_csv(z / 'electrical' / 'wire_list.csv')
    wires_csv = {r['Wire'] for r in wl if r['Wire']}
    wires_xl = {str(r['Wire']) for r in sheet_rows(wb, 'Wire List') if r.get('Wire')}
    wires_lbl = {l['Number'] for l in labels if l['Type'] == 'Wire markers'}
    if wires_csv != wires_xl:
        report('FAIL', v, f'wire numbers CSV vs Excel differ: {sorted(wires_csv ^ wires_xl)[:8]}')
    if wires_csv != wires_lbl:
        report('FAIL', v, f'wire numbers CSV vs labels differ: {sorted(wires_csv ^ wires_lbl)[:8]}')
    # wire markers: one per conductor end (2 per wire-list row of that number)
    ends = {}
    for r in wl:
        if r['Wire']: ends[r['Wire']] = ends.get(r['Wire'], 0) + 2
    short = [f"{l['Number']} {l['Copies']}/{ends.get(l['Number'])}" for l in labels if l['Type'] == 'Wire markers' and int(l['Copies']) != ends.get(l['Number'])]
    if short:
        report('FAIL', v, f'wire marker copies differ from 2 per conductor: {short[:8]}')
    # pypdf joins texts on one baseline without a space ("Y003AOL2" / IEC "Y004A-F1" = wire, then contact tag)
    missing = sorted(w for w in wires_csv if not re.search(r'(?<![\w-])' + re.escape(w) + r'(?:(?![\w-])|(?=-?[A-Z]{1,4}\d))', drw_all))
    if missing:
        report('FAIL', v, f'{len(missing)} wire numbers from the wire list never appear on the drawings: {missing[:12]}')
    for r in wl:
        for k in ('From', 'To'):
            if re.search(r'\?|undefined|null', r[k]):
                report('WARN', v, f'wire {r["Wire"]} {k} = "{r[k]}" (unassigned device)')
    # marker copies = 2 per conductor
    per = {}
    for r in wl:
        per[r['Wire']] = per.get(r['Wire'], 0) + 1
    for l in labels:
        if l['Type'] == 'Wire markers' and int(l['Copies']) != 2 * per.get(l['Number'], 0):
            report('FAIL', v, f'wire marker {l["Number"]}: {l["Copies"]} copies, wire list has {per.get(l["Number"], 0)} conductors')

    # ---- Terminals
    tp = read_csv(z / 'electrical' / 'terminal_plan.csv')
    t_csv = {(f'{r["Strip"]}-{r["Terminal"]}', r['Wire']) for r in tp}
    t_xl = {(str(r['Terminal']), str(r['Wire'] or '')) for r in sheet_rows(wb, 'Terminal Plan') if r.get('Terminal') and r.get('Internal (panel)') is not None or (r.get('Terminal') and r.get('Kind'))}
    if t_csv != t_xl:
        report('FAIL', v, f'terminal plan CSV vs Excel differ: {sorted(t_csv ^ t_xl)[:6]}')
    # terminal markers "TB2-4" / IEC "-X2-4" (not the strip group markers "TB2" / "-X2")
    t_lbl = {l['Number'] for l in labels if l['Type'] == 'Terminal markers' and re.search(r'.-\d+$', l['Number'])}
    if {t for t, _ in t_csv} != t_lbl:
        report('FAIL', v, f'terminal labels vs plan differ: {sorted({t for t, _ in t_csv} ^ t_lbl)[:8]}')
    # strip size vs BOM quantity of feed-through terminals
    n_term = len(tp)
    tb_q = sum(q for pn, q in csv_q.items() if pn.startswith(('DN-T', 'DN-F')))  # feed-through + fused
    if tb_q and tb_q < n_term - sum(1 for r in tp if r.get('Type') == 'ground'):
        report('WARN', v, f'BOM has {tb_q} feed-through terminals, terminal plan needs {n_term}')

    # ---- I/O addresses
    io_csv = read_csv(z / 'io_list.csv')
    io_a = {(r['PLC Address'], r['Description']) for r in io_csv}
    io_x = {(str(r['Address']), str(r['Description'])) for r in sheet_rows(wb, 'I-O List')}
    io_d = {(x['address'], x['description']) for x in data['io']}
    if io_a != io_x or io_a != io_d:
        report('FAIL', v, f'I/O list CSV / Excel / app differ: {sorted(io_a ^ io_x)[:4]} {sorted(io_a ^ io_d)[:4]}')
    nick = next(z.glob('plc/*_nicknames.csv'), None)
    nick_txt = nick.read_text(encoding='utf-8-sig') if nick else ''
    # CLICK analog channels (DF registers) are not BIT rows: they go in the analog setup CSV instead.
    ana = next(z.glob('plc/*_analog_setup.csv'), None)
    ana_txt = ana.read_text(encoding='utf-8-sig') if ana else ''
    for a, desc in io_a:
        if not a:
            report('FAIL', v, f'I/O "{desc}" has no PLC address')
            continue
        if not re.search(r'(?<!\w)' + a + r'(?!\w)', drw_all):
            report('FAIL', v, f'I/O address {a} not on the drawings')
        if a not in eng_all:
            report('FAIL', v, f'I/O address {a} not in the engineering package')
        if re.fullmatch(r'DF\d+', a):
            if not re.search(r'(?<!\w)' + a + r'(?!\w)', ana_txt):
                report('FAIL', v, f'analog I/O address {a} not in the analog setup CSV')
        elif nick and a not in nick_txt:
            report('FAIL', v, f'I/O address {a} not in the nickname CSV')

    # ---- Alarms
    hmi_al = [r for r in sheet_rows(wb, 'HMI Alarms')] if 'HMI Alarms' in wb.sheetnames else []
    for a in data['alarms']:
        msg = a['message']
        if msg not in eng_flat:
            report('FAIL', v, f'alarm "{msg}" not in the engineering package')
        if hmi_al and not any(str(r.get('Message')) == msg and str(r.get('Address')) == a['address'] for r in hmi_al):
            report('FAIL', v, f'alarm {a["address"]} "{msg}" differs between alarm list and HMI Alarms sheet')

    # ---- Revision letter / document number on every page
    rev = data['revision']['rev']
    for i, t in enumerate(eng):
        if f'Rev {rev}' not in t:
            report('FAIL', v, f'engineering package p{i + 1}: no "Rev {rev}" in header/footer')
    for i, t in enumerate(drw):
        tb = flat(t)
        if not re.search(r'\bREV\b', tb) or not re.search(r'(?<![A-Z])' + rev + r'(?![A-Z])', tb):
            report('FAIL', v, f'drawing sheet {i + 1}: revision {rev} not in the title block')
    dc = [r for r in wb['Document Control'].iter_rows(values_only=True) if r[0] == 'Current revision']
    if not dc or not str(dc[0][1]).startswith(f'{rev} '):
        report('FAIL', v, f'Excel document control revision {dc[0][1] if dc else None} vs {rev}')

    # ---- Motor FLA (independent table) and calculations
    for mo in data['motors']:
        ref = nec_fla(mo['hp'], mo['voltage'])
        if abs(mo['fla'] - ref) > 1e-6:
            report('FAIL', v, f'{mo["tag"]} {mo["hp"]} HP {mo["voltage"]}: FLA {mo["fla"]} A, NEC/CEC table says {ref} A')
        if f'{mo["fla"]} A' not in eng_flat:
            report('WARN', v, f'{mo["tag"]} FLA {mo["fla"]} A not printed in the engineering package')
    # ---- Cabinet cooling: heat the walls cannot shed (rise above the limit) needs a filter fan, a heat exchanger or an
    # air conditioner in the design, or the design review saying so; a filter fan comes with its exhaust filter and
    # thermostat, and a 24 VDC fan is wired (its wire on the drawings and in the wire list)
    co = data.get('cooling') or {}
    kinds = {c['kind'] for c in co.get('parts', [])}
    if co.get('ventilation') and not kinds & {'filter fan', 'heat exchanger', 'air conditioner'} and not any('Cabinet cooling missing' in w for w in data['warnings']):
        report('FAIL', v, f'cabinet needs {co["ventilation"]["airflowM3h"]:.0f} m³/h of cooling ({co["rise"]:.0f} °C rise) but has no fan / heat exchanger / air conditioner and the review does not say so')
    if 'filter fan' in kinds:
        for need in ('exhaust filter', 'thermostat'):
            if need not in kinds:
                report('WARN', v, f'filter fan without a {need}')
        fan = co.get('fan') or {}
        if fan.get('dc') and fan.get('thermostat') and fan['tag'] not in {w['wire'] for w in data['wires']}:
            report('FAIL', v, f'24 VDC fan {fan["tag"]}: no wire {fan["tag"]} from thermostat {fan["thermostat"]} in the wire list')
    # ---- Drive branches: the drive's listed Class T fuses (its SCCR condition), or the review says why not
    iec = str(data.get('region', '')) not in ('US', 'Canada', '')
    for dr in data.get('drives', []):
        if str(dr.get('fuseClass') or '').upper() != 'T' or not dr.get('fusePart') or iec:
            continue
        pr = dr.get('protection') or {}
        fine = pr.get('category') == 'fuse' and str(pr.get('fuseClass') or '').upper() == 'T' and str(pr.get('fuse') or '') in (dr['fusePart'], f'{dr["fusePart"]}-1')
        flagged = any(dr['tag'] in w and 'Class T' in w for w in data['warnings'])
        if not fine and not flagged:
            report('FAIL', v, f'{dr["tag"]} {dr["partNumber"]}: branch protection {pr.get("tag")} {pr.get("partNumber")} is not its Class T {dr["fusePart"]} fuses, and the design review does not say why')
    for s in data['calcs']:
        if s['title'].startswith('24 VDC'):
            tot = sum(float(l['value'].split()[0]) for l in s['lines'])
            m = re.search(r'Load ([\d.]+) A \+ 25% margin = ([\d.]+) A; supply rated ([\d.]+) A', s['result'])
            if not m or abs(float(m.group(1)) - tot) > 0.005 * len(s['lines']) + 0.006 or abs(float(m.group(2)) - float(m.group(1)) * 1.25) > 0.011:
                report('FAIL', v, f'24 VDC load does not add up: lines {tot:.2f} A vs "{s["result"]}"')
            elif (float(m.group(3)) >= float(m.group(2))) != s['ok']:
                report('FAIL', v, '24 VDC OK flag inconsistent')
        if s['title'].startswith('Cabinet'):
            tot = sum(float(l['value'].split()[0]) for l in s['lines'] if l['value'].endswith(' W'))
            m = re.match(r'(\d+) W in ([\d.]+) m', s['result'])
            if not m or abs(int(m.group(1)) - tot) > 0.6:
                report('FAIL', v, f'heat does not add up: lines {tot:.1f} W vs "{s["result"][:30]}"')
            dr = [str(r[0]) for r in wb['Design Review'].iter_rows(values_only=True)] if 'Design Review' in wb.sheetnames else []
            if s['ok'] is False and not any('temperature rise' in w.lower() for w in dr):
                report('FAIL', v, f'heat check fails ({s["result"][:60]}…) but the Excel design review list does not mention it')
        if s['title'].startswith('Motor') and not s['ok']:
            report('WARN', v, f'{s["title"]}: {s["result"]}')
    print(f'{v}: BOM {len(csv_q)} parts, {len(wires_csv)} wire numbers, {len(tp)} terminals, {len(io_a)} I/O, {len(data["alarms"])} alarms, rev {rev}')


def main():
    names = sys.argv[1:] or sorted(p.name for p in OUT.iterdir() if (p / 'data.json').exists())
    for n in names:
        check(OUT / n)
    for lvl, where, msg in results:
        print(f'{lvl:4} {where}: {msg}')
    fails = [r for r in results if r[0] == 'FAIL']
    print(f'CONSISTENCY: {len(fails)} FAIL, {len(results) - len(fails)} WARN')
    sys.exit(1 if fails else 0)


if __name__ == '__main__':
    main()
